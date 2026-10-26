"""The two original files, for a manual comparison side by side.

One workbook, one sheet per file: the original columns and rows as RIM-Soft
exported them, without the matched operations, and a coloured «الحالة» column:
red (on this side only), orange (possibly the same operation), yellow (a
different amount), grey (outside the chosen period). The original files are
never modified: only a new file is written, where the user chooses.
"""

from __future__ import annotations

from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from binaa_reconciler.compare.engine import (
    Kind,
    Line,
    Pair,
    Result,
    amount_text,
)
from binaa_reconciler.compare.files import Loaded
from binaa_reconciler.compare.texts import reason
from binaa_reconciler.core.models import LedgerSide

RED = PatternFill("solid", fgColor="F8CBAD")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="E7E6E6")
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
RLM = "\u200f"  # keeps Latin names and numbers in place inside Arabic text


ARABIC_NAMES = {"binaa": "بناء", "clinique": "كلينك", "clinic": "كلينك", "cliniq": "كلينك"}


def arabic(name: str) -> str:
    """The party's name in Arabic letters when known: no Latin word inside an
    Arabic sentence, so the sentence never breaks (else the name as typed)."""
    return ARABIC_NAMES.get(name.strip().casefold(), name)


def _arabic_text(text: str, names: tuple[str, str]) -> str:
    for name in names:
        text = text.replace(name, arabic(name))
    return text


def _rtl(text: str) -> str:
    """Right-to-left text: one fact per line, each line (and each piece
    between the separators) anchored right, so names, dates and amounts stay
    where they are read."""
    lines = text.replace(" — ", "\n").split("\n")
    return "\n".join(RLM + line.replace(": ", f":{RLM} ").replace(" ", f" {RLM}") for line in lines)


DATE_FORMAT = "DD/MM/YYYY"
_FORBIDDEN = str.maketrans(dict.fromkeys("[]:*?/\\", " "))


def _other(pair: Pair, line: Line) -> str:
    """The other side of a pair, as «dd/mm amount (+ ...)»."""
    side = pair.b if line in pair.a else pair.a
    if not side:  # a reversal in the same file
        side = tuple(x for x in pair.lines if x is not line)
    return " + ".join(f"{x.day:%d/%m} {amount_text(x.amount)}" for x in side)


def statuses(result: Result, i: int) -> tuple[set[int], dict[int, tuple[str, PatternFill]]]:
    """Rows to drop (matched) and the status of each remaining row of file i."""
    other = result.names[1 - i]
    matched: set[int] = set()
    for x, y in result.matched:
        matched.add((x, y)[i].row)
    matched.update(x.row for x in result.settled if x.key[0] == "AB"[i])
    for pair in result.confirmed:
        matched.update(x.row for x in (pair.a if i == 0 else pair.b))
        if pair.kind is Kind.REVERSAL:
            matched.update(x.row for x in pair.lines if x.key[0] == "AB"[i])
    status: dict[int, tuple[str, PatternFill]] = {}
    for pair in result.possible:
        for x in pair.lines:
            if x.key[0] == "AB"[i]:
                text = f"قد يكون: {reason(pair, result.names)} — المقابل: {_other(pair, x)}"
                status[x.row] = (text, ORANGE)
    for pair in result.close:
        for x in pair.lines:
            if x.key[0] == "AB"[i]:
                text = f"المقابل عند {other}: {_other(pair, x)}"  # yellow: a gap
                status[x.row] = (text, YELLOW)
    for x in result.only[i]:
        note = result.notes.get(x.key)
        text = f"عند {result.names[i]} فقط" + (f" — {note}" if note else "")
        status[x.row] = (text, RED)
    status = {row: (_arabic_text(t, result.names), f) for row, (t, f) in status.items()}
    return matched, status


def _owed(side: LedgerSide, other: str) -> str:
    """In a party's own statement: debit = owed by the other party, credit =
    owed to it."""
    return f"على {other}" if side is LedgerSide.DEBIT else f"ل{other}"


def _copy_sheet(target: Worksheet, loaded: Loaded, result: Result, i: int) -> None:
    source_wb = load_workbook(BytesIO(loaded.data), data_only=True)
    source = source_wb.worksheets[loaded.sheet_index]
    matched, status = statuses(result, i)
    other = arabic(result.names[1 - i])
    lines = {x.row: x for x in loaded.lines}
    second_rows = {v: k for k, v in loaded.continuation.items()}  # 2nd row -> 1st
    rows = list(source.iter_rows())
    width = source.max_column
    status_col = width + 1

    # the remaining operations, grouped: what is owed by the other party, what
    # is owed to it (each alone on this side), then the amount gaps, the
    # possible pairs and what is outside the period; the file order inside
    groups: dict[str, list[int]] = {}
    for n in range(loaded.header_row + 1, len(rows) + 1):
        x = lines.get(n)
        if x is None or n in matched:
            continue  # matched, a 2nd row (follows its operation) or a total line
        fill = status.get(n, ("", GREY))[1]
        if fill is RED:
            title = f"{_owed(x.side, other)} — عند {arabic(result.names[i])} فقط"
        else:
            title = {YELLOW: "فرق في المبلغ", ORANGE: "قد تكون نفس العملية"}.get(
                fill, "خارج الفترة المختارة"
            )
        groups.setdefault(title, []).append(n)
    order = [f"على {other}", f"ل{other}", "فرق", "قد تكون", "خارج"]
    titles = sorted(groups, key=lambda t: next(k for k, o in enumerate(order) if t.startswith(o)))

    out = 0

    def copy_row(n: int) -> None:
        nonlocal out
        out += 1
        for cell in rows[n - 1]:
            value: Any = cell.value
            new = target.cell(out, cell.column or 1, value)
            if cell.has_style:
                new.number_format = cell.number_format
                new.font = copy(cell.font)
            if isinstance(value, datetime | date):
                new.number_format = DATE_FORMAT
        first = second_rows.get(n, n)
        if n in lines and loaded.date_col is not None:  # never a bare number
            target.cell(out, loaded.date_col + 1, lines[n].day).number_format = DATE_FORMAT
        if n > loaded.header_row:  # amounts: 1,216,000; the empty side stays empty
            for col in loaded.amount_cols:
                cell = target.cell(out, col + 1)
                if isinstance(cell.value, int | float):
                    cell.value = cell.value or None
                    cell.number_format = "#,##0"
            text, fill = status.get(first, ("خارج الفترة المختارة", GREY))
            if n == first:
                target.cell(out, status_col, _rtl(text)).alignment = WRAP
            for col in range(1, status_col + 1):
                target.cell(out, col).fill = fill

    for n in range(1, loaded.header_row + 1):
        copy_row(n)
    target.cell(out, status_col, "الحالة").font = Font(bold=True)
    for title in titles:
        members = groups[title]
        total = sum(lines[n].amount for n in members)
        out += 1
        head = f"{title} — {len(members)} عملية — المجموع {amount_text(total)}"
        target.cell(out, 1, RLM + head).font = Font(bold=True, size=12)
        for n in members:
            copy_row(n)
            if (second := loaded.continuation.get(n)) is not None:
                copy_row(second)
    for col in range(1, width + 1):
        letter = get_column_letter(col)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width or 14
    # narrow note columns, the text wraps: the page fits the screen
    target.column_dimensions[get_column_letter(status_col)].width = 42
    target.freeze_panes = f"A{loaded.header_row + 1}"


def build(result: Result, a: Loaded, b: Loaded) -> bytes:
    wb = Workbook()
    if wb.active is not None:
        wb.remove(wb.active)
    for i, loaded in enumerate((a, b)):
        title = arabic(result.names[i]).translate(_FORBIDDEN)[:31] or "AB"[i]
        ws: Worksheet = wb.create_sheet(title)
        _copy_sheet(ws, loaded, result, i)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save(result: Result, a: Loaded, b: Loaded, target: Path) -> Path:
    target.write_bytes(build(result, a, b))
    return target


__all__ = ["build", "save", "statuses"]
