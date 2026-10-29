"""Sale prices check: the product lines copied from RIM-Soft (Statistiques →
Mouvements), pasted into a spreadsheet and saved (xlsx / csv / txt).

Columns used: Code, Désignation, Vendu (quantity) and PrixVente. For each
product the usual price is learned (this file and the files checked before,
kept on this computer); a price far from it is pointed out with its line
number in the file, so it is found at once in RIM-Soft. Nothing is changed.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path

GAP = 0.25  # a price 25 % away from the product's usual one
LEARNED = 3  # prices of a product needed before it is judged


@dataclass(frozen=True)
class PriceLine:
    row: int  # line number in the file (Excel row)
    code: str
    name: str
    quantity: float
    price: float


@dataclass(frozen=True)
class PriceFinding:
    line: PriceLine
    usual: float
    text: str


@dataclass
class PriceHistory:
    prices: dict[str, list[float]] = field(default_factory=dict)

    def learn(self, lines: Iterable[PriceLine]) -> None:
        for x in lines:
            if x.price > 0:
                self.prices.setdefault(x.code, []).append(x.price)
                del self.prices[x.code][:-200]

    def to_json(self) -> dict[str, list[float]]:
        return self.prices

    @staticmethod
    def from_json(data: object) -> PriceHistory:
        h = PriceHistory()
        if isinstance(data, dict):
            try:
                h.prices = {str(k): [float(p) for p in v] for k, v in data.items()}
            except (TypeError, ValueError):
                return PriceHistory()
        return h


def _number(value: object) -> float | None:
    if isinstance(value, int | float):
        return float(value)
    text = str(value or "").replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _rows(path: Path) -> list[list[object]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    dialect = "excel-tab" if "\t" in text.splitlines()[0] else "excel"
    if dialect == "excel" and ";" in text.splitlines()[0]:
        return [list(r) for r in csv.reader(io.StringIO(text), delimiter=";")]
    return [list(r) for r in csv.reader(io.StringIO(text), dialect=dialect)]


def read_lines(path: Path) -> list[PriceLine]:
    rows = _rows(path)
    found = [
        (h, [str(c or "").strip().lower() for c in header])
        for h, header in enumerate(rows[:10])
        if {"code", "prixvente"} <= {str(c or "").strip().lower() for c in header}
    ]
    if not found:
        raise ValueError("no_header")
    h, names = found[0]
    col = {n: k for k, n in enumerate(names)}
    name_col = next((k for n, k in col.items() if n.startswith("d")), col["code"])
    qty_col = col.get("vendu")
    out = []
    for n, row in enumerate(rows[h + 1 :], start=h + 2):
        code = str(row[col["code"]] or "").strip() if len(row) > col["code"] else ""
        price = _number(row[col["prixvente"]]) if len(row) > col["prixvente"] else None
        if not code or price is None:
            continue
        qty = _number(row[qty_col]) if qty_col is not None and len(row) > qty_col else 0
        out.append(PriceLine(n, code, str(row[name_col] or "").strip(), qty or 0.0, price))
    return out


def check_prices(
    lines: Sequence[PriceLine], history: PriceHistory | None = None
) -> list[PriceFinding]:
    known = PriceHistory.from_json(
        {k: list(v) for k, v in (history or PriceHistory()).prices.items()}
    )
    known.learn(lines)
    found = []
    for x in lines:
        others = list(known.prices.get(x.code, []))
        if x.price in others:
            others.remove(x.price)
        if len(others) < LEARNED:
            continue
        usual = sorted(others)[len(others) // 2]
        if usual <= 0:
            continue
        if x.price <= 0:
            text = f"سعر صفر — المعتاد {usual:,.0f}"
        elif abs(x.price - usual) / usual > GAP:
            ratio = x.price / usual
            how = (
                f"{ratio:.0f} مرات أكبر"
                if ratio >= 2
                else (
                    f"{1 / ratio:.0f} مرات أصغر"
                    if ratio <= 0.5
                    else f"{abs(ratio - 1):.0%} {'أكبر' if ratio > 1 else 'أصغر'}"
                )
            )
            text = f"السعر {x.price:,.0f} والمعتاد {usual:,.0f} — {how}"
        else:
            continue
        found.append(PriceFinding(x, usual, text))
    return found


__all__ = ["PriceFinding", "PriceHistory", "PriceLine", "check_prices", "read_lines"]
