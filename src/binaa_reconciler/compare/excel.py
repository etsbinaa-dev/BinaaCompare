"""The comparison as an Excel file: one sheet per list, right-to-left, whole
amounts as real numbers."""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from binaa_reconciler.compare.engine import SIDE_AR, Line, Pair, Result, units
from binaa_reconciler.compare.texts import footnote, headline, reason, side_cells, title

MONEY = "#,##0"
_HEAD = PatternFill("solid", fgColor="E8EEF4")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
_FORBIDDEN = str.maketrans(dict.fromkeys("[]:*?/\\", " "))


def _safe(text: str) -> str:
    """Never a formula: a cell text starting with = + - @ is written as text."""
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def _sheet(wb: Workbook, name: str, head: Sequence[str], widths: Sequence[int]) -> Worksheet:
    ws: Worksheet = wb.create_sheet(name.translate(_FORBIDDEN)[:31])
    ws.sheet_view.rightToLeft = True
    ws.append(list(head))
    for k, width in enumerate(widths, start=1):
        cell = ws.cell(1, k)
        cell.font, cell.fill, cell.border = Font(bold=True), _HEAD, _BOX
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(k)].width = width
    ws.freeze_panes = "A2"
    return ws


def _finish(ws: Worksheet, money: Sequence[int]) -> None:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = _BOX
            if cell.column in money and isinstance(cell.value, int):
                cell.number_format = MONEY
    ws.auto_filter.ref = ws.dimensions


def _alone(wb: Workbook, result: Result, i: int) -> None:
    name = result.names[i]
    ws = _sheet(
        wb,
        f"عند {name} فقط",
        ("التاريخ", "المبلغ", "الجهة", "الوصف", "المرجع", "ملاحظة"),
        (12, 14, 8, 44, 16, 40),
    )
    lines: list[Line] = result.only[i]
    for x in lines:
        ws.append(
            [
                f"{x.day:%d/%m/%Y}",
                units(x.amount),
                SIDE_AR[x.side],
                _safe(x.label),
                _safe(x.ref),
                _safe(result.notes.get(x.key, "")),
            ]
        )
    ws.append(["المجموع", units(result.total(i))])
    ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = Font(bold=True)
    _finish(ws, (2,))


def _pairs(wb: Workbook, result: Result, name: str, pairs: Sequence[Pair]) -> None:
    a, b = result.names
    ws = _sheet(
        wb,
        name,
        (f"تاريخ {a}", f"مبلغ {a}", f"وصف {a}", f"تاريخ {b}", f"مبلغ {b}", f"وصف {b}",
         "الفرق", "السبب"),
        (12, 14, 30, 12, 14, 30, 12, 40),
    )  # fmt: skip
    for pair in pairs:
        da, ma, la = side_cells(pair.a)
        db, mb, lb = side_cells(pair.b)
        ws.append(
            [
                da,
                units(sum(x.amount for x in pair.a)) if pair.a else ma,
                _safe(la),
                db,
                units(sum(x.amount for x in pair.b)) if pair.b else mb,
                _safe(lb),
                units(pair.difference) if pair.a and pair.b else "",
                _safe(reason(pair, result.names)),
            ]
        )
    _finish(ws, (2, 5, 7))


def build(result: Result) -> bytes:
    wb = Workbook()
    if wb.active is not None:
        wb.remove(wb.active)
    summary = wb.create_sheet("الملخص")
    summary.sheet_view.rightToLeft = True
    summary.column_dimensions["A"].width = 110
    for text, bold in ((title(result), True), (headline(result), True), (footnote(result), False)):
        if text:
            summary.append([_safe(text)])
            summary.cell(summary.max_row, 1).font = Font(bold=bold)
    _pairs(wb, result, "قد تكون نفس العملية", result.possible)
    _pairs(wb, result, "فرق في المبلغ", result.close)
    _alone(wb, result, 0)
    _alone(wb, result, 1)
    matched = [*(((x,), (y,)) for x, y in result.matched), *((p.a, p.b) for p in result.confirmed)]
    a, b = result.names
    ws = _sheet(
        wb,
        "المتطابقة",
        (f"تاريخ {a}", f"مبلغ {a}", f"تاريخ {b}", f"مبلغ {b}", f"وصف {a}", f"وصف {b}"),
        (12, 14, 12, 14, 34, 34),
    )
    for left, right in sorted(matched, key=lambda m: min(x.day for x in (*m[0], *m[1]))):
        da, ma, la = side_cells(left)
        db, mb, lb = side_cells(right)
        ws.append(
            [
                da,
                units(sum(x.amount for x in left)) if left else ma,
                db,
                units(sum(x.amount for x in right)) if right else mb,
                _safe(la),
                _safe(lb),
            ]
        )
    _finish(ws, (2, 4))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save(result: Result, target: Path) -> Path:
    target.write_bytes(build(result))
    return target


__all__ = ["build", "save"]
