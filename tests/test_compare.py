"""BinaaCompare (synthetic data only): two RIM-Soft files of one account."""

from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook

from binaa_reconciler.compare import for_party, manual, memory
from binaa_reconciler.compare.engine import Line, amount_text, compare
from binaa_reconciler.compare.excel import build
from binaa_reconciler.compare.files import load
from binaa_reconciler.compare.texts import footnote, reason
from binaa_reconciler.core.models import LedgerSide

D, C = LedgerSide.DEBIT, LedgerSide.CREDIT
NAMES = ("binaa", "clinique")


def d(n: int) -> date:
    return date(2026, 9, 1) + timedelta(days=n - 1)


def line(key: str, day: int, side: LedgerSide, amount: float, label: str = "SYN") -> Line:
    return Line(key, int(key[1:]), d(day), side, round(amount * 100), label)


EDGES = (
    [line("A901", 1, D, 1), line("A902", 30, D, 2)],
    [line("B901", 1, C, 1), line("B902", 30, C, 2)],
)  # matched lines on the 1st and the 30th: both files cover the whole month


def run(a, b):
    return compare([*a, *EDGES[0]], [*b, *EDGES[1]], NAMES)


def run_with(a, b, **decisions):
    return compare([*a, *EDGES[0]], [*b, *EDGES[1]], NAMES, None, **decisions)


def keys(pairs):
    return [(tuple(x.key for x in p.a), tuple(x.key for x in p.b), p.kind.value) for p in pairs]


def test_mirror_lines_up_to_three_days_apart_match():
    r = run(
        [line("A1", 2, D, 1000), line("A2", 5, C, 500), line("A3", 9, D, 70)],
        [line("B1", 2, C, 1000), line("B2", 8, D, 500), line("B3", 9, D, 70)],
    )
    pairs = [(x.key, y.key) for x, y in r.matched if x.key[1] != "9"]
    assert pairs == [("A1", "B1"), ("A2", "B2")]
    # same side on both files: never matched, a possible pair to check
    assert keys(r.possible) == [(("A3",), ("B3",), "SIDE")]
    assert r.only == ([], [])


def test_possible_pairs_are_a_separate_list_close_amounts_another():
    r = run(
        [
            line("A1", 2, D, 1000),
            line("A2", 3, D, 10_000),
            line("A3", 4, D, 3_000),
            line("A4", 6, D, 700),
            line("A5", 7, C, 700),
            line("A6", 12, D, 444),
        ],
        [
            line("B1", 9, C, 1000),
            line("B2", 4, C, 10_150),
            line("B3", 4, C, 1_200),
            line("B4", 4, C, 1_800),
            line("B5", 20, C, 555),
        ],
    )
    assert keys(r.possible) == [
        (("A4", "A5"), (), "REVERSAL"),
        (("A3",), ("B3", "B4"), "SUM"),
        (("A1",), ("B1",), "SAME"),
    ]
    assert keys(r.close) == [(("A2",), ("B2",), "CLOSE")]
    assert r.close[0].difference == -15_000
    # alone: nothing on the other side could be it
    assert [x.key for x in r.only[0]] == ["A6"] and [x.key for x in r.only[1]] == ["B5"]
    assert reason(r.possible[2], NAMES) == "نفس المبلغ — فرق 7 يوم في التاريخ"
    assert reason(r.close[0], NAMES) == "فرق في المبلغ: 150"


def test_confirm_makes_a_pair_disappear_reject_frees_its_lines():
    a, b = [line("A1", 2, D, 1000)], [line("B1", 9, C, 1000)]
    r = run(a, b)
    (pair,) = r.possible
    ok = run_with(a, b, confirmed=[pair])
    assert ok.possible == [] and ok.only == ([], []) and ok.confirmed == [pair]
    assert ok.matched_count == 3  # the two edge lines + the confirmed pair
    no = run_with(a, b, rejected={pair.key})
    assert no.possible == [] and [x.key for x in no.only[0]] == ["A1"]
    assert [x.key for x in no.only[1]] == ["B1"]


def test_duplicates_pair_as_far_as_possible_the_extra_is_listed():
    r = run(
        [line("A1", 5, D, 2000), line("A2", 5, D, 2000), line("A3", 6, D, 2000)],
        [line("B1", 5, C, 2000), line("B2", 6, C, 2000)],
    )
    assert len(r.matched) == 4 and [x.key for x in r.only[0]] == ["A2"]


def test_the_period_lists_its_lines_and_counts_the_others():
    a = [line("A1", 1, D, 10), line("A2", 10, D, 20), line("A3", 25, D, 30)]
    b = [line("B1", 5, C, 40), line("B2", 20, C, 50)]
    r = compare(a, b, NAMES)  # default: the days both files cover
    assert r.period == (d(5), d(20))
    assert [x.key for x in r.only[0]] == ["A2"] and r.outside == (2, 0)
    assert footnote(r).startswith("غير معروضة لأنها خارج الفترة المختارة: 2 عملية عند binaa")
    wide = compare(a, b, NAMES, (d(1), d(30)))
    assert [x.key for x in wide.only[0]] == ["A1", "A2", "A3"] and wide.outside == (0, 0)
    assert wide.notes["A1"] == "قبل أول عملية في ملف clinique (05/09/2026)"
    assert wide.notes["A3"] == "بعد آخر عملية في ملف clinique (20/09/2026)"
    assert footnote(wide) == ""


def test_a_counterpart_just_outside_the_period_still_matches():
    r = compare([line("A1", 20, D, 500)], [line("B1", 22, C, 500)], NAMES, (d(1), d(20)))
    assert len(r.matched) == 1 and r.only == ([], [])


def test_amounts_are_whole_units_one_block_of_digits():
    assert amount_text(414_000_000) == "4,140,000" and amount_text(12_350) == "124"
    assert amount_text(-15_000) == "-150"


def test_excel_one_sheet_per_list_whole_numbers():
    r = compare(
        [line("A1", 2, D, 1000), line("A2", 3, D, 4140000.5), line("A3", 4, D, 10_000)],
        [line("B1", 2, C, 1000), line("B2", 5, C, 10_100)],
        NAMES,
        (d(1), d(30)),
    )
    wb = load_workbook(io.BytesIO(build(r)))
    assert wb.sheetnames == [
        "الملخص",
        "قد تكون نفس العملية",
        "فرق في المبلغ",
        "عند binaa فقط",
        "عند clinique فقط",
        "المتطابقة",
    ]
    alone = list(wb["عند binaa فقط"].iter_rows(min_row=2, values_only=True))
    assert (
        alone[0][:3] == ("03/09/2026", 4140001, "مدين")
        and wb["عند binaa فقط"]["B2"].number_format == "#,##0"
    )
    close = list(wb["فرق في المبلغ"].iter_rows(min_row=2, values_only=True))
    assert close[0][1] == 10_000 and close[0][4] == 10_100 and close[0][6] == -100
    assert next(iter(wb["المتطابقة"].iter_rows(min_row=2, values_only=True)))[:4] == (
        "02/09/2026",
        1000,
        "02/09/2026",
        1000,
    )


def _xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.create_sheet("Feuil1", 0)
    ws.append(["N°", "Date", "Compte", "Libellé", "Pièce", "Débit", "Crédit", "Solde"])
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_real_shaped_file_zero_rows_and_a_split_row_are_read():
    data = _xlsx(
        [
            (1, date(2026, 9, 14), 900, "SYN TRT", "OD-1", 0, 1500, 0),
            (2, date(2026, 9, 14), 900, "SYN BIM"),  # cut by a line break in the label
            (None, "OD-2", 0, 1870, 0, 0),
            (3, date(2026, 9, 15), 900, "SYN info", "OD-3", 0, 0, 0),  # zero: skipped
            (4, date(2026, 9, 16), 900, "SYN sans montant", "OD-4"),  # cannot be rejoined
        ]
    )
    loaded = load(data, "A")
    assert loaded.error is None
    assert [x.amount for x in loaded.lines] == [150_000, 187_000]
    assert loaded.merged == 1 and len(loaded.problems) == 1
    assert loaded.problems[0].startswith("السطر 6:")


def test_an_unreadable_file_is_a_message_not_a_crash():
    assert load(b"not an excel file", "A").error is not None


def test_manual_export_drops_matched_rows_and_groups_what_is_owed():
    a = load(
        _xlsx(
            [
                (1, date(2026, 9, 14), 900, "SYN paye", "OD-1", 5000, 0, 0),  # matched
                (2, date(2026, 9, 14), 900, "SYN BIM"),  # split: two rows, kept together
                (None, "OD-2", 0, 1870, 0, 0),
                (3, date(2026, 9, 15), 900, "SYN seul", "OD-3", 0, 777, 0),
                (4, date(2026, 9, 16), 900, "SYN livre", "OD-4", 900, 0, 0),
            ]
        ),
        "A",
    )
    b = load(_xlsx([(1, date(2026, 9, 14), 411, "SYN recu", "X-1", 0, 5000, 0)]), "B")
    result = compare(a.lines, b.lines, ("binaa", "clinique"), (d(1), d(30)))
    wb = load_workbook(io.BytesIO(manual.build(result, a, b)))
    assert wb.sheetnames == ["بناء", "كلينك"]  # Arabic names: no broken sentence
    rows = [tuple(r) for r in wb["بناء"].iter_rows(values_only=True)]
    assert rows[0][:3] == ("N°", "Date", "Compte") and rows[0][8] == "الحالة"
    plain = [[str(c).replace("\u200f", "") for c in r] for r in rows]
    # owed by clinique first (debit), then owed to it (credit); row 1 (matched) gone
    assert plain[1][0] == "على كلينك — عند بناء فقط — 1 عملية — المجموع 900"
    assert rows[2][0] == 4
    assert plain[3][0] == "لكلينك — عند بناء فقط — 2 عملية — المجموع 2,647"
    assert [r[0] for r in rows[4:]] == [2, None, 3]
    assert plain[4][8] == "عند بناء فقط"
    assert rows[5][8] is None and rows[5][1] == "OD-2"  # the 2nd row, with its operation
    assert plain[6][8].startswith("عند بناء فقط\nبعد آخر عملية في ملف كلينك")
    assert len(rows[0]) == 9  # one status column
    assert rows[6][5:7] == (None, 777)  # the empty side (0 in RIM-Soft) stays empty
    assert wb["بناء"]["A3"].fill.fgColor.rgb.endswith("F8CBAD")
    assert [r[0] for r in wb["كلينك"].iter_rows(values_only=True)] == ["N°"]


def test_five_transfers_of_one_day_making_one_entry_of_the_other_file():
    """Bankily transfers recorded one by one on one side, as one total (a day
    later) on the other; an unrelated line of that day stays alone."""
    parts = [
        line(f"B{i}", 5, C, amount) for i, amount in enumerate((1_200, 3_500, 800, 2_000, 4_500), 1)
    ]
    r = run([line("A1", 6, D, 12_000)], [*parts, line("B6", 5, C, 950)])
    assert keys(r.possible) == [(("A1",), ("B1", "B2", "B3", "B4", "B5"), "SUM")]
    assert reason(r.possible[0], NAMES) == "مجموع 5 عمليات من يوم واحد يساوي عملية واحدة"
    assert [x.key for x in r.only[1]] == ["B6"]


def test_a_sum_needs_lines_of_one_day_and_never_stalls_on_a_busy_day():
    split = [line("B1", 5, C, 5_000), line("B2", 6, C, 7_000)]  # two different days
    assert keys(run([line("A1", 6, D, 12_000)], split).possible) == []
    busy = [line(f"B{i}", 5, C, 1_000 + i) for i in range(1, 200)]
    r = run([line("A1", 5, D, 7)], busy)  # nothing can make 7: answered at once
    assert r.possible == []


def test_saved_decisions_follow_the_lines_not_the_excel_rows():
    a = [line("A3", 4, D, 3_000)]
    b = [line("B3", 4, C, 1_200), line("B4", 4, C, 1_800), line("B5", 9, C, 77)]
    (pair,) = run(a, b).possible
    memory.save(NAMES, [(True, pair)])
    moved = [line("B13", 4, C, 1_200), line("B14", 4, C, 1_800)]  # exported again
    ((ok, again),) = memory.load(NAMES, a, moved)
    assert ok and [x.key for x in again.b] == ["B13", "B14"]
    assert memory.load(NAMES, a, moved[:1]) == []  # a line is gone: dropped
    assert memory.load(("x", "y"), a, moved) == []  # other parties: their own


def test_manual_export_writes_dates_even_when_the_file_has_serial_numbers():
    a = load(_xlsx([(1, 46279, 900, "SYN serial", "OD-1", 0, 777, 0)]), "A")  # 14/09/2026
    b = load(_xlsx([(1, date(2026, 9, 20), 411, "SYN autre", "X-1", 0, 5, 0)]), "B")
    result = compare(a.lines, b.lines, ("binaa", "clinique"), (d(1), d(30)))
    ws = load_workbook(io.BytesIO(manual.build(result, a, b)))["بناء"]
    assert ws["B3"].value.date() == date(2026, 9, 14) and ws["B3"].number_format == "DD/MM/YYYY"
    assert ws["I3"].alignment.wrap_text and ws["I3"].alignment.readingOrder == 2


def test_the_same_amount_alone_on_both_sides_is_pointed_out():
    r = run([line("A1", 2, D, 4_000)], [line("B1", 25, C, 4_000)])  # 23 days apart
    assert r.notes["A1"] == "⚠️ نفس المبلغ عند clinique بتاريخ 25/09/2026"
    assert r.notes["B1"] == "⚠️ نفس المبلغ عند binaa بتاريخ 02/09/2026"


def _xlsx4(rows) -> bytes:
    """The short RIM-Soft layout: Date, Libellé, Débit, Crédit only."""
    wb = Workbook()
    ws = wb.create_sheet("Feuil1", 0)
    ws.append(["Date", "Libellé", "Débit", "Crédit"])
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_split_rows_are_rejoined_in_the_short_layout_too():
    before = [(date(2026, 9, 13), "SYN avant", 500, 0)]
    after = [(date(2026, 9, 15), "SYN apres", 0, 900)]
    for second in (
        (None, 1870, 0),  # the cut-off cells, from the 2nd column
        (1870, 0),  # the cut-off cells, from the 1st column
        ("suite", 1870, 0),  # the end of the label, then the cells
    ):
        loaded = load(_xlsx4([*before, (date(2026, 9, 14), "SYN BIM"), second, *after]), "B")
        assert [x.amount for x in loaded.lines] == [50_000, 187_000, 90_000], second
        assert loaded.merged == 1 and loaded.problems == (), second


def test_a_row_without_amount_before_a_real_operation_is_not_merged():
    loaded = load(
        _xlsx4([(date(2026, 9, 14), "SYN sans montant"), (date(2026, 9, 15), "SYN ok", 0, 900)]),
        "B",
    )
    assert [x.amount for x in loaded.lines] == [90_000] and loaded.merged == 0
    assert len(loaded.problems) == 1


def test_one_short_file_for_both_parties_by_date():
    r = run(
        [line("A1", 3, D, 500), line("A2", 4, C, 200), line("A3", 6, D, 10_000)],
        [
            line("B1", 7, C, 10_150),
            line("B2", 8, D, 70),
            line("B3", 9, C, 30),
            line("B4", 15, C, 500),  # the same amount 12 days later: a possible pair
            line("B5", 12, C, 1_200),
            line("B6", 12, C, 800),
        ],
    )
    ws = load_workbook(io.BytesIO(for_party.build(r))).active
    assert ws.title == "بناء وكلينك" and ws.sheet_view.rightToLeft
    values = [tuple(r) for r in ws.iter_rows(values_only=True)]
    head = values.index(("التاريخ", "البيان", "مبلغ بناء", "مبلغ كلينك", "ملاحظة"))
    body = [r for r in values[head + 1 :] if r[0]]
    assert body[:-1] == [
        ("04/09/2026", "SYN", 200, "ناقصة", "ناقصة عند كلينك — لصالح كلينك"),
        # the other date and description only where they differ
        (
            "06/09/2026",
            "SYN",
            10_000,
            10_150,
            "فرق في المبلغ: 150 — عند كلينك بتاريخ 07/09/2026 — بيان كلينك: SYN",
        ),
        ("08/09/2026", "SYN", "ناقصة", 70, "ناقصة عند بناء — لصالح كلينك"),
        ("09/09/2026", "SYN", "ناقصة", 30, "ناقصة عند بناء — لصالح بناء"),
        ("12/09/2026", "SYN", "ناقصة", 800, "ناقصة عند بناء — لصالح بناء"),
        ("12/09/2026", "SYN", "ناقصة", 1_200, "ناقصة عند بناء — لصالح بناء"),
    ]  # the possible pair (500 / 500) is not there: decided in the application
    assert body[-1][0] == "مجموع الفروق في الأسطر الصفراء (1 عملية): 150"  # the only total
    assert "A1:E1" in {str(m) for m in ws.merged_cells.ranges}  # the title is never cut
    first = head + 2
    assert ws.cell(first, 4).value == "ناقصة"  # the missing side, marked and darker
    assert ws.cell(first, 4).fill.fgColor.rgb.endswith("FF8B8B")
    assert ws.cell(first, 1).fill.fgColor.rgb.endswith("F8CBAD")  # the whole line is red


def test_a_line_settled_by_hand_leaves_the_list_and_is_remembered():
    from binaa_reconciler.compare.engine import Kind, Pair
    from binaa_reconciler.compare.texts import headline

    a, b = [line("A1", 3, D, 500), line("A2", 4, C, 200)], [line("B1", 8, D, 70)]
    settled = [(True, Pair(Kind.SETTLED, (a[0],), ())), (True, Pair(Kind.SETTLED, (), (b[0],)))]
    r = run_with(a, b, confirmed=[p for _, p in settled])
    assert [x.key for x in r.only[0]] == ["A2"] and r.only[1] == []
    assert [x.key for x in r.settled] == ["A1", "B1"] and r.matched_count == 2  # edges only
    assert "✔️ تمت تسويتها: 2" in headline(r)
    memory.save(NAMES, settled)
    assert [p.kind for _, p in memory.load(NAMES, a, b)] == [Kind.SETTLED, Kind.SETTLED]


def test_why_a_line_is_alone_although_the_other_side_has_the_same_amount():
    from binaa_reconciler.compare.texts import same_amount

    a = [line("A1", 20, D, 3_402)]
    b = [line("B1", 21, C, 3_402), line("B2", 5, C, 3_402), line("B3", 9, C, 77)]
    r = run(a, b)
    (alone,) = [x for x in r.only[1] if x.amount == 340_200]  # B2: A1 went to B1
    assert same_amount(r, alone, a, "binaa") == (
        "نفس المبلغ عند binaa: 20/09/2026 مدين: مطابقة مع clinique 21/09"
    )
    (other,) = [x for x in r.only[1] if x.key == "B3"]
    assert same_amount(r, other, a, "binaa") == "لا توجد عند binaa أي عملية بالمبلغ 77"


def test_a_delivery_note_number_written_twice_in_one_file_is_pointed_out():
    a = [
        line("A1", 14, D, 3_402, "60t chinguitt BL 38957 BR 26A00587"),
        line("A2", 20, D, 3_402, "60t chinguitt bl38957"),  # the same delivery, again
        line("A3", 22, D, 900, "NFT 26T01660"),  # NFT: given by RIM-Soft, never compared
        line("A4", 23, D, 700, "NFT 26T01660"),
    ]
    r = run(a, [line("B1", 14, C, 3_402)])
    assert r.notes["A2"].startswith("⚠️ رقم BL 38957 مسجّل مرتين في ملف binaa (14/09/2026)")
    assert "مرتين" not in r.notes.get("A3", "") + r.notes.get("A4", "")
