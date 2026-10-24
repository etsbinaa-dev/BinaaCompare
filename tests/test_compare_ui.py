"""BinaaCompare window (synthetic files only)."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

pytest.importorskip("tkinter")

from binaa_reconciler.compare.ui import CompareApp


def xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["N°", "Date", "Compte", "Libellé", "Pièce", "Débit", "Crédit", "Solde"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


class FakeDialogs:
    def __init__(self, files, save_to):
        self.files, self.save_to, self.calls = list(files), save_to, []

    def ask_open(self):
        return self.files.pop(0)

    def ask_save(self, suggested):
        return self.save_to

    def info(self, title, text):
        self.calls.append(("info", title))

    def error(self, title, text):
        self.calls.append(("error", title))

    def open_file(self, path):
        self.calls.append(("open", path.name))


@pytest.fixture
def files(tmp_path):
    a = xlsx(
        tmp_path / "a.xlsx",
        [
            (1, date(2026, 9, 2), 900, "SYN Fact 1", "T1", 1000, 0, 0),
            (2, date(2026, 9, 3), 900, "SYN Fact 2", "T2", 500, 0, 0),
            (3, date(2026, 9, 9), 900, "SYN Vers.", "OD1", 0, 300, 0),
        ],
    )
    b = xlsx(
        tmp_path / "b.xlsx",
        [
            (1, date(2026, 9, 2), 900, "SYN Fact 1", "T1", 0, 1000, 0),
            (2, date(2026, 9, 9), 900, "SYN Paiement", "P1", 0, 300, 0),
        ],
    )
    return a, b


def test_two_files_one_button_two_lists(files, tmp_path):
    try:
        app = CompareApp(FakeDialogs(files, tmp_path / "out.xlsx"))
    except Exception as exc:  # no display
        pytest.skip(f"Tk cannot open a window here: {type(exc).__name__}")
    try:
        app.update()
        assert app.compare_button.instate(["disabled"])
        app.box_a.choose()
        app.box_b.choose()
        assert app.box_a.status_var.get().startswith("✅ a.xlsx: 3 عملية")
        assert (app.from_var.get(), app.to_var.get()) == ("02/09/2026", "09/09/2026")
        app.compare_button.invoke()
        app.update()
        head = app.headline_var.get()
        assert "✅ متطابقة: 1" in head and "🔎 قد تكون نفس العملية: 1" in head
        alone = [app.trees[0].item(k, "values")[::-1] for k in app.trees[0].get_children()]
        assert [row[1] for row in alone] == ["500"]  # whole amounts, no decimals
        # 300 credit on both sides: a possible pair (entry side to check), not alone
        possible = app.pair_trees["possible"]
        (key,) = possible.get_children()
        assert "تحقق من جهة القيد" in possible.item(key, "values")[0]
        # ✓ مطابقة: the pair disappears and joins the matches
        possible.selection_set(key)
        app.decide("possible", True)
        app.update()
        assert possible.get_children() == ()
        assert "✅ متطابقة: 2" in app.headline_var.get()
        # ↶ undo: the pair is proposed again, then confirmed again
        app.undo()
        app.update()
        assert possible.get_children() == (key,)
        possible.selection_set(key)
        app.decide("possible", True)
        app.update()
        # a shorter period: the lines outside it are counted, not listed
        app.from_var.set("09/09/2026")
        app.compare_button.invoke()
        app.update()
        assert app.trees[0].get_children() == ()
        assert "غير معروضة لأنها خارج الفترة المختارة: 1 عملية عند binaa" in app.note_var.get()
        target = app.export_manual()
        assert target is not None and target.exists()
        assert load_workbook(io.BytesIO(target.read_bytes())).sheetnames == ["binaa", "clinique"]
        assert app.dialogs.calls[-1] == ("open", "out.xlsx")  # opened after saving
    finally:
        app.destroy()


def test_the_same_file_twice_is_refused(files, tmp_path):
    try:
        app = CompareApp(FakeDialogs([files[0], files[0]], tmp_path / "out.xlsx"))
    except Exception as exc:  # no display
        pytest.skip(f"Tk cannot open a window here: {type(exc).__name__}")
    try:
        app.box_a.choose()
        app.box_b.choose()
        app.compare_button.invoke()
        assert app.result is None and app.dialogs.calls == [("error", "الملف نفسه")]
    finally:
        app.destroy()


def test_decisions_survive_closing_the_window(files, tmp_path):
    """A crash or a power cut: the ✓ / ✗ are back when the files are compared again."""
    for step in ("decide", "reopen"):
        try:
            app = CompareApp(FakeDialogs(files, tmp_path / "out.xlsx"))
        except Exception as exc:  # no display
            pytest.skip(f"Tk cannot open a window here: {type(exc).__name__}")
        try:
            app.box_a.choose()
            app.box_b.choose()
            app.compare_button.invoke()
            app.update()
            possible = app.pair_trees["possible"]
            if step == "decide":
                (key,) = possible.get_children()
                possible.selection_set(key)
                app.decide("possible", True)
            else:
                assert possible.get_children() == ()
                assert "✅ متطابقة: 2" in app.headline_var.get()
                assert "استُرجع 1 قرارًا محفوظًا" in app.note_var.get()
                app.undo()  # undo is saved too
                assert len(possible.get_children()) == 1
        finally:
            app.destroy()
